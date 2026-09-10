import json
import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

from src.data_store import DataStore
from src.review_decision import ReviewDecisionService
from src.runtime import file_sha256, read_json, write_json
from src.sampling_export import (
    SAMPLING_DISCLAIMER,
    SamplingExportService,
    SamplingExportValidationError,
    SamplingListEmptyError,
    write_sampling_workbook,
)
from src.sampling_store import SamplingStore
from tests.test_data_store import create_run


def write_recommendation(product_root: Path, label: str = "助眠相关宣传线索") -> None:
    write_json(
        product_root / "inspection_context.json",
        {
            "product_category": "固体饮料",
            "product_form": "粉剂",
            "confirmed_ingredient_contexts": [],
            "context_evidence": [],
        },
    )
    write_json(
        product_root / "inspection_recommendation.json",
        {
            "product_context": {
                "product_category": "固体饮料",
                "product_form": "粉剂",
                "confirmed_ingredient_contexts": [],
                "context_evidence": [],
            },
            "risk_findings": [
                {
                    "risk_category": "sleep",
                    "risk_labels": [label],
                    "possible_risk_summary": "页面中发现助眠相关宣传线索。",
                    "evidence_qualification": "seller_managed_primary",
                    "substance_follow_ups": [
                        {
                            "substance_id": "substance-1",
                            "canonical_name": "测试关注成分",
                            "english_name": "",
                            "cas_no": "",
                            "regulatory_context_note": "",
                            "follow_up_status": "suggest_testing",
                            "reason": "基于已保存页面线索供人工判断。",
                            "suggested_methods": [
                                {
                                    "method_id": "method-1",
                                    "method_no": "BJS TEST 001",
                                    "method_name": "测试方法",
                                    "method_type": "screening",
                                    "method_status": "current",
                                    "determination_role": "screening",
                                    "applicability_status": "applicable",
                                    "applicability_reason": "测试适用范围匹配",
                                    "source_name": "测试已核验来源",
                                    "source_reference": "https://example.invalid/reference",
                                    "source_date": "2026-01-01",
                                }
                            ],
                            "methods_needing_context": [],
                            "other_known_methods": [],
                        }
                    ],
                }
            ],
            "unmapped_evidence": [],
            "composition_gaps": [],
            "knowledge_gaps": [],
            "disclaimer": "测试冻结免责声明",
        },
    )


class SamplingExportServiceTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.output = self.root / "output"
        self.run_root = create_run(
            self.output,
            "run-a",
            product_id="123",
            product_name="=危险公式商品",
            display_name="九月抽检排查",
        )
        product_root = self.run_root / "products" / "123"
        (product_root / "ocr").mkdir(parents=True, exist_ok=True)
        (product_root / "ocr" / "original_001.txt").write_text(
            "帮助睡眠", encoding="utf-8"
        )
        (product_root / "page").mkdir(parents=True, exist_ok=True)
        (product_root / "page" / "overview.png").write_bytes(b"frozen-image")
        write_recommendation(product_root)

        self.data_store = DataStore(self.root / "data" / "app.db", self.output)
        self.data_store.initialize()
        self.data_store.import_all_runs()
        self.sampling_store = SamplingStore(self.data_store.database_path)
        self.decisions = ReviewDecisionService(
            self.data_store, self.sampling_store
        )
        self.snapshot_id = self.data_store.list_products()[0]["snapshotId"]
        self.service = SamplingExportService(
            self.data_store, self.sampling_store, self.output
        )

    def tearDown(self):
        self.temporary.cleanup()

    def add_current(self, note: str = "人工确认") -> None:
        self.decisions.decide(
            self.snapshot_id,
            "recommend_follow_up",
            "inspection_workspace",
            note,
        )

    def test_export_freezes_json_xlsx_and_clears_only_membership(self):
        self.add_current("+人工备注公式")
        metadata = self.service.export_current(confirmed=True)

        self.assertEqual(metadata["status"], "exported")
        self.assertEqual(metadata["itemCount"], 1)
        self.assertEqual(self.sampling_store.count_current(), 0)
        snapshot = self.data_store.get_snapshot(self.snapshot_id)
        self.assertEqual(snapshot["review"]["status"], "recommend_follow_up")
        self.assertEqual(snapshot["review"]["note"], "+人工备注公式")
        self.assertEqual(snapshot["sampling"]["decisionStatus"], "reviewed_follow_up")

        list_root = self.output / "sampling_lists" / metadata["listId"]
        snapshot_path = list_root / "sampling_list_snapshot.json"
        workbook_path = list_root / "sampling_list.xlsx"
        payload = read_json(snapshot_path)
        self.assertEqual(payload["items"][0]["sourceSnapshotId"], self.snapshot_id)
        self.assertEqual(payload["items"][0]["summary"]["substances"], ["测试关注成分"])
        self.assertTrue(payload["items"][0]["frozenAssets"])
        self.assertTrue(
            all(
                (list_root / asset["frozenPath"]).is_file()
                and file_sha256(list_root / asset["frozenPath"]) == asset["sha256"]
                for asset in payload["items"][0]["frozenAssets"]
            )
        )
        self.assertEqual(file_sha256(snapshot_path), metadata["snapshotSha256"])
        self.assertEqual(file_sha256(workbook_path), metadata["workbookSha256"])

        workbook = load_workbook(workbook_path, data_only=False)
        self.assertEqual(workbook.sheetnames, ["抽检辅助清单", "说明"])
        sheet = workbook["抽检辅助清单"]
        self.assertEqual(sheet["A2"].value, "'=危险公式商品")
        self.assertNotEqual(sheet["A2"].data_type, "f")
        self.assertEqual(sheet["M2"].value, "'+人工备注公式")
        self.assertEqual(sheet["B2"].hyperlink.target, "https://item.taobao.com/item.htm?id=123")
        self.assertEqual(workbook["说明"]["B7"].value, SAMPLING_DISCLAIMER)

        with sqlite3.connect(self.data_store.database_path) as connection:
            index_row = connection.execute(
                "SELECT product_id, source_snapshot_id, source_task_id "
                "FROM sampling_list_item_index WHERE list_id = ?",
                (metadata["listId"],),
            ).fetchone()
        self.assertEqual(index_row, ("123", self.snapshot_id, "run-a"))

    def test_export_failure_preserves_membership_and_has_no_success_history(self):
        self.add_current()
        failing = SamplingExportService(
            self.data_store,
            self.sampling_store,
            self.output,
            workbook_writer=Mock(side_effect=RuntimeError("disk unavailable")),
        )
        with self.assertRaisesRegex(Exception, "disk unavailable"):
            failing.export_current(confirmed=True)
        self.assertEqual(self.sampling_store.count_current(), 1)
        self.assertEqual(self.sampling_store.list_history(status="exported"), [])
        self.assertEqual(self.sampling_store.list_history(status="preparing"), [])

    def test_preparing_export_is_recovered_after_process_interruption(self):
        self.add_current()
        self.service._finalize_export = Mock(side_effect=SystemExit("crash"))
        with self.assertRaises(SystemExit):
            self.service.export_current(confirmed=True)
        self.assertEqual(self.sampling_store.count_current(), 1)
        self.assertEqual(len(self.sampling_store.list_history(status="preparing")), 1)

        restarted = SamplingExportService(
            self.data_store, self.sampling_store, self.output
        )
        result = restarted.recover_preparing()
        self.assertEqual(result, {"recovered": 1, "cleaned": 0})
        self.assertEqual(self.sampling_store.count_current(), 0)
        self.assertEqual(len(restarted.list_history()["lists"]), 1)

    def test_incomplete_preparing_export_is_cleaned_without_touching_membership(self):
        self.add_current()
        list_id = "SL-INCOMPLETE-TEST"
        with self.data_store.transaction() as connection:
            self.sampling_store.create_preparing(
                connection,
                list_id=list_id,
                snapshot_path=f"sampling_lists/{list_id}/sampling_list_snapshot.json",
                workbook_path=f"sampling_lists/{list_id}/sampling_list.xlsx",
                created_at="2026-09-10T10:00:00+08:00",
            )
        partial_root = self.output / "sampling_lists" / list_id
        partial_root.mkdir(parents=True)
        write_json(partial_root / "sampling_list_snapshot.json", {"incomplete": True})

        result = self.service.recover_preparing()
        self.assertEqual(result, {"recovered": 0, "cleaned": 1})
        self.assertEqual(self.sampling_store.count_current(), 1)
        self.assertFalse(partial_root.exists())
        self.assertIsNone(self.sampling_store.get_history_metadata(list_id))

    def test_export_does_not_clear_membership_recreated_during_file_write(self):
        self.add_current()

        def replace_membership_then_write(items, list_id, exported_at, destination):
            self.sampling_store.remove("123")
            self.decisions.restore_membership(
                "123", self.snapshot_id, "inspection_workspace"
            )
            write_sampling_workbook(
                items, list_id, exported_at, destination
            )

        service = SamplingExportService(
            self.data_store,
            self.sampling_store,
            self.output,
            workbook_writer=replace_membership_then_write,
        )
        metadata = service.export_current(confirmed=True)
        self.assertEqual(metadata["status"], "exported")
        self.assertEqual(self.sampling_store.count_current(), 1)
        self.assertEqual(
            self.sampling_store.get("123")["addedFrom"], "inspection_workspace"
        )

    def test_frozen_detail_does_not_drift_after_live_files_and_review_change(self):
        self.add_current()
        metadata = self.service.export_current(confirmed=True)
        before = self.service.get_history(metadata["listId"])

        write_recommendation(
            self.run_root / "products" / "123", "后来改变的风险方向"
        )
        self.data_store.update_review(
            self.snapshot_id, "no_further_action", "后来改变的备注"
        )
        after = self.service.get_history(metadata["listId"])
        self.assertEqual(after, before)
        self.assertEqual(
            after["items"][0]["summary"]["riskDirections"],
            ["助眠相关宣传线索"],
        )

    def test_old_run_without_recommendation_exports_without_inventing_methods(self):
        product_root = self.run_root / "products" / "123"
        (product_root / "inspection_recommendation.json").unlink()
        (product_root / "inspection_context.json").unlink()
        self.add_current()

        metadata = self.service.export_current(confirmed=True)
        item = self.service.get_history(metadata["listId"])["items"][0]
        self.assertEqual(item["summary"]["substances"], [])
        self.assertEqual(item["summary"]["methods"], [])
        self.assertEqual(item["summary"]["riskDirections"], ["助眠"])
        self.assertEqual(item["disclaimer"], SAMPLING_DISCLAIMER)

    def test_history_index_rebuilds_without_original_run_or_source_entities(self):
        self.add_current()
        metadata = self.service.export_current(confirmed=True)
        new_database = self.root / "rebuilt" / "app.db"
        shutil.rmtree(self.run_root)

        rebuilt_data = DataStore(new_database, self.output)
        rebuilt_data.initialize()
        rebuilt_store = SamplingStore(new_database)
        rebuilt_service = SamplingExportService(
            rebuilt_data, rebuilt_store, self.output
        )
        result = rebuilt_service.rebuild_history_index()
        self.assertEqual(result, {"rebuilt": 1, "skipped": 0})
        self.assertEqual(
            rebuilt_service.get_history(metadata["listId"])["items"][0]["productId"],
            "123",
        )
        with sqlite3.connect(new_database) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM tasks").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM products").fetchone()[0], 0)
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM sampling_list_item_index"
                ).fetchone()[0],
                1,
            )

    def test_empty_unconfirmed_and_invalid_history_ids_are_rejected(self):
        with self.assertRaises(SamplingExportValidationError):
            self.service.export_current(confirmed=False)
        with self.assertRaises(SamplingListEmptyError):
            self.service.export_current(confirmed=True)
        with self.assertRaises(SamplingExportValidationError):
            self.service.get_history("../outside")

    def test_workbook_rejects_formula_cells_and_non_http_product_links(self):
        destination = self.root / "unsafe.xlsx"
        write_sampling_workbook(
            [
                {
                    "productName": "@unsafe",
                    "productUrl": "javascript:alert(1)",
                    "shopName": "-unsafe",
                    "sourceTaskDisplayName": "任务",
                    "collectedAt": "2026-09-10T10:00:00+08:00",
                    "summary": {
                        "riskDirections": [],
                        "evidenceQualifications": [],
                        "substances": [],
                        "methods": [],
                    },
                    "evidence": [],
                    "review": {"note": "=1+1"},
                }
            ],
            "SL-TEST-1",
            "2026-09-10T10:00:00+08:00",
            destination,
        )
        workbook = load_workbook(destination, data_only=False)
        row = workbook["抽检辅助清单"]
        self.assertEqual(row["A2"].value, "'@unsafe")
        self.assertIsNone(row["B2"].value)
        self.assertIsNone(row["B2"].hyperlink)
        self.assertEqual(row["C2"].value, "'-unsafe")
        self.assertEqual(row["M2"].value, "'=1+1")
        self.assertTrue(all(row.cell(2, column).data_type != "f" for column in range(1, 14)))


if __name__ == "__main__":
    unittest.main()
