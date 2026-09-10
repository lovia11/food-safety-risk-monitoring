"""Freeze the current sampling list into immutable JSON and XLSX facts."""

from __future__ import annotations

import json
import os
import re
import shutil
import threading
import uuid
from copy import deepcopy
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.data_store import DataStore
from src.runtime import file_sha256, iso_now, read_json, write_json
from src.sampling_store import SamplingStore
from src.web_contract import build_snapshot_artifacts


SAMPLING_SNAPSHOT_FILE = "sampling_list_snapshot.json"
SAMPLING_WORKBOOK_FILE = "sampling_list.xlsx"
SAMPLING_SNAPSHOT_VERSION = 1
SAMPLING_DISCLAIMER = (
    "本清单仅用于监管抽检辅助筛查和人工研判，不构成实验室检测结论，"
    "不认定商品违法、功效真实、实际含有或检出任何成分。具体检验项目及方法"
    "适用性应由检验人员结合样品身份、类别、剂型、配料和现行标准确认。"
)

_LIST_ID_PATTERN = re.compile(r"^SL-[A-Za-z0-9][A-Za-z0-9-]{0,63}$")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_EXPORT_LOCKS: dict[Path, threading.Lock] = {}
_EXPORT_LOCKS_GUARD = threading.Lock()


class SamplingExportError(RuntimeError):
    """Base error for sampling-list export and history reads."""


class SamplingExportValidationError(SamplingExportError):
    """The export request or list identifier is invalid."""


class SamplingListEmptyError(SamplingExportError):
    """There is no current Membership to freeze."""


class SamplingExportConflictError(SamplingExportError):
    """Current data cannot be exported without violating Review semantics."""


class SamplingExportInProgressError(SamplingExportError):
    """Another export currently owns the process lock."""


class SamplingHistoryNotFoundError(SamplingExportError):
    """The requested frozen list does not exist."""


class SamplingHistoryIntegrityError(SamplingExportError):
    """A frozen fact no longer matches its recorded identity or hash."""


def validate_list_id(list_id: str) -> str:
    value = str(list_id or "").strip()
    if not _LIST_ID_PATTERN.fullmatch(value):
        raise SamplingExportValidationError("历史清单编号不合法")
    return value


def _safe_http_url(value: Any) -> str:
    text = str(value or "").strip()
    try:
        parsed = urlparse(text)
    except ValueError:
        return ""
    return text if parsed.scheme in {"http", "https"} and parsed.netloc else ""


def _safe_cell(value: Any, *, limit: int = 32000) -> str:
    text = str(value or "")[:limit]
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _unique_text(values: list[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _sampling_summary(
    snapshot: dict[str, Any], inspection: dict[str, Any]
) -> dict[str, Any]:
    risk_directions: list[Any] = []
    qualifications: list[Any] = []
    substances: list[Any] = []
    methods: list[dict[str, str]] = []
    method_ids: set[str] = set()
    for finding in inspection.get("riskFindings") or []:
        labels = finding.get("risk_labels") or []
        risk_directions.extend(labels or [finding.get("risk_category")])
        qualifications.append(finding.get("evidence_qualification"))
        for substance in finding.get("substance_follow_ups") or []:
            substances.append(substance.get("canonical_name"))
            for group in (
                "suggested_methods",
                "methods_needing_context",
                "other_known_methods",
            ):
                for method in substance.get(group) or []:
                    identity = str(
                        method.get("method_id")
                        or method.get("method_no")
                        or method.get("method_name")
                        or ""
                    )
                    if not identity or identity in method_ids:
                        continue
                    method_ids.add(identity)
                    methods.append(
                        {
                            "methodId": str(method.get("method_id") or ""),
                            "methodNo": str(method.get("method_no") or ""),
                            "methodName": str(method.get("method_name") or ""),
                            "applicabilityStatus": str(
                                method.get("applicability_status") or ""
                            ),
                            "applicabilityReason": str(
                                method.get("applicability_reason") or ""
                            ),
                            "sourceName": str(method.get("source_name") or ""),
                            "sourceReference": str(
                                method.get("source_reference") or ""
                            ),
                        }
                    )
    if not risk_directions:
        risk_directions.extend(snapshot.get("detectedEffects") or [])
    return {
        "riskDirections": _unique_text(risk_directions),
        "evidenceQualifications": _unique_text(qualifications),
        "substances": _unique_text(substances),
        "methods": methods,
    }


def _major_evidence(evidence: list[dict[str, Any]]) -> list[str]:
    seller = [
        item.get("text")
        for item in evidence
        if item.get("contentOrigin") == "seller_managed"
    ]
    return _unique_text(seller or [item.get("text") for item in evidence])


def _matches_asset_path(candidate: Any, source_path: Any) -> bool:
    left = str(candidate or "").replace("\\", "/").lstrip("/").lower()
    right = str(source_path or "").replace("\\", "/").lstrip("/").lower()
    return bool(left and right and (left == right or left.endswith("/" + right)))


def _asset_candidates(
    product_id: str,
    evidence: list[dict[str, Any]],
    assets: dict[str, Any],
) -> list[tuple[str, str]]:
    candidates: list[tuple[str, str]] = []
    for item in evidence:
        source_path = item.get("sourcePath")
        if source_path:
            candidates.append(
                ("evidence_source", f"products/{product_id}/{source_path}")
            )
        for ocr_item in assets.get("ocrItems") or []:
            if not any(
                _matches_asset_path(ocr_item.get(field), source_path)
                for field in ("textPath", "jsonPath")
            ):
                continue
            for field, kind in (
                ("imagePath", "evidence_image"),
                ("textPath", "ocr_text"),
                ("jsonPath", "ocr_json"),
            ):
                if ocr_item.get(field):
                    candidates.append((kind, str(ocr_item[field])))
    if assets.get("overview"):
        candidates.append(("page_overview", str(assets["overview"])))
    result: list[tuple[str, str]] = []
    seen: set[str] = set()
    for kind, path in candidates:
        normalized = path.replace("\\", "/").lstrip("/")
        if normalized and normalized not in seen:
            seen.add(normalized)
            result.append((kind, normalized))
    return result


def write_sampling_workbook(
    items: list[dict[str, Any]], list_id: str, exported_at: str, destination: Path
) -> None:
    """Write the user-facing workbook without formulas or inferred facts."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "抽检辅助清单"
    sheet.sheet_view.showGridLines = False
    headers = [
        "商品名称",
        "商品链接",
        "店铺",
        "来源排查",
        "页面采集时间",
        "可能风险方向",
        "主要页面证据",
        "Evidence qualification",
        "建议关注/检测成分",
        "相关方法编号",
        "方法名称",
        "方法适用状态",
        "人工备注",
    ]
    sheet.append(headers)
    for item in items:
        summary = item.get("summary") or {}
        methods = summary.get("methods") or []
        row = [
            _safe_cell(item.get("productName")),
            _safe_cell(_safe_http_url(item.get("productUrl"))),
            _safe_cell(item.get("shopName")),
            _safe_cell(item.get("sourceTaskDisplayName") or item.get("sourceTaskId")),
            _safe_cell(item.get("collectedAt")),
            _safe_cell("\n".join(summary.get("riskDirections") or [])),
            _safe_cell("\n".join(_major_evidence(item.get("evidence") or []))),
            _safe_cell("\n".join(summary.get("evidenceQualifications") or [])),
            _safe_cell("\n".join(summary.get("substances") or [])),
            _safe_cell("\n".join(_unique_text([m.get("methodNo") for m in methods]))),
            _safe_cell("\n".join(_unique_text([m.get("methodName") for m in methods]))),
            _safe_cell(
                "\n".join(
                    _unique_text([m.get("applicabilityStatus") for m in methods])
                )
            ),
            _safe_cell((item.get("review") or {}).get("note")),
        ]
        sheet.append(row)
        url = _safe_http_url(item.get("productUrl"))
        if url:
            sheet.cell(sheet.max_row, 2).hyperlink = url
            sheet.cell(sheet.max_row, 2).style = "Hyperlink"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="1F2937")
    thin = Side(style="thin", color="D9E1E8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    table = Table(displayName="SamplingAuxiliaryList", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    widths = [28, 34, 20, 22, 23, 24, 48, 24, 24, 20, 30, 22, 34]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.row_dimensions[1].height = 30

    notes = workbook.create_sheet("说明")
    notes.sheet_view.showGridLines = False
    note_rows = [
        ["抽检辅助清单说明", ""],
        ["清单编号", list_id],
        ["导出时间", exported_at],
        ["商品数量", len(items)],
        ["字段说明", "风险方向、成分与方法均来自导出时已保存的页面 Evidence、人工 Review 与抽检辅助建议。"],
        ["Evidence qualification", "seller_managed_primary 表示主要线索来自商家管理内容；user_generated_auxiliary_only 仅表示用户生成内容辅助线索。"],
        ["免责声明", SAMPLING_DISCLAIMER],
    ]
    for row in note_rows:
        notes.append([_safe_cell(value) for value in row])
    notes["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color="1F2937")
    for row in notes.iter_rows(min_row=2):
        row[0].font = Font(name="Microsoft YaHei", size=10, bold=True, color="334155")
        row[1].font = body_font
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 88
    notes.row_dimensions[5].height = 42
    notes.row_dimensions[6].height = 54
    notes.row_dimensions[7].height = 72
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


class SamplingExportService:
    """Coordinate current DB state, frozen files, history index and export recovery."""

    def __init__(
        self,
        data_store: DataStore,
        sampling_store: SamplingStore,
        output_root: Path,
        *,
        workbook_writer: Callable[[list[dict[str, Any]], str, str, Path], None]
        | None = None,
    ) -> None:
        self.data_store = data_store
        self.sampling_store = sampling_store
        self.output_root = output_root.resolve()
        self.sampling_root = (self.output_root / "sampling_lists").resolve()
        self.workbook_writer = workbook_writer or write_sampling_workbook
        with _EXPORT_LOCKS_GUARD:
            self._export_lock = _EXPORT_LOCKS.setdefault(
                self.data_store.database_path, threading.Lock()
            )

    def _list_root(self, list_id: str) -> Path:
        validated = validate_list_id(list_id)
        destination = (self.sampling_root / validated).resolve()
        if destination.parent != self.sampling_root:
            raise SamplingExportValidationError("历史清单路径不合法")
        return destination

    def _relative_paths(self, list_id: str) -> tuple[str, str]:
        base = f"sampling_lists/{validate_list_id(list_id)}"
        return f"{base}/{SAMPLING_SNAPSHOT_FILE}", f"{base}/{SAMPLING_WORKBOOK_FILE}"

    def _run_root(self, snapshot: dict[str, Any]) -> Path:
        relative = str((snapshot.get("paths") or {}).get("run") or "").strip()
        destination = (self.output_root / relative).resolve()
        if not relative or destination.parent != self.output_root:
            raise SamplingExportConflictError("商品快照的来源排查路径不合法")
        return destination

    def _copy_frozen_assets(
        self,
        *,
        run_root: Path,
        product_id: str,
        item_ordinal: int,
        evidence: list[dict[str, Any]],
        assets: dict[str, Any],
        temporary_assets_root: Path,
    ) -> list[dict[str, Any]]:
        frozen: list[dict[str, Any]] = []
        for asset_ordinal, (kind, relative) in enumerate(
            _asset_candidates(product_id, evidence, assets), start=1
        ):
            source = (run_root / relative).resolve()
            if not source.is_relative_to(run_root) or not source.is_file():
                continue
            safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", source.name) or "asset"
            frozen_relative = (
                Path("assets")
                / f"item-{item_ordinal:04d}"
                / f"{asset_ordinal:03d}-{safe_name}"
            )
            temporary_destination = temporary_assets_root / Path(*frozen_relative.parts[1:])
            temporary_destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source, temporary_destination)
            frozen.append(
                {
                    "kind": kind,
                    "sourcePath": relative,
                    "frozenPath": frozen_relative.as_posix(),
                    "sha256": file_sha256(temporary_destination),
                }
            )
        return frozen

    def _item_from_snapshot(
        self,
        membership: dict[str, Any],
        snapshot: dict[str, Any],
        artifacts: dict[str, Any],
        *,
        ordinal: int,
        temporary_assets_root: Path | None = None,
    ) -> dict[str, Any]:
        evidence = deepcopy(snapshot.get("evidence") or [])
        inspection = deepcopy(artifacts.get("inspection") or {})
        run_root = self._run_root(snapshot)
        frozen_assets = (
            self._copy_frozen_assets(
                run_root=run_root,
                product_id=str(snapshot.get("productId") or ""),
                item_ordinal=ordinal,
                evidence=evidence,
                assets=artifacts.get("assets") or {},
                temporary_assets_root=temporary_assets_root,
            )
            if temporary_assets_root is not None
            else []
        )
        return {
            "ordinal": ordinal,
            "productId": membership["productId"],
            "sourceSnapshotId": membership["sourceSnapshotId"],
            "sourceTaskId": membership["sourceTaskId"],
            "sourceTaskDisplayName": snapshot.get("taskDisplayName") or membership["sourceTaskId"],
            "addedFrom": membership["addedFrom"],
            "addedAt": membership["addedAt"],
            "updatedAt": membership["updatedAt"],
            "productName": snapshot.get("productName") or "",
            "productUrl": snapshot.get("productUrl") or "",
            "shopName": snapshot.get("shopName") or "",
            "collectedAt": snapshot.get("collectedAt"),
            "detectedEffects": deepcopy(snapshot.get("detectedEffects") or []),
            "historicalCountBeforeExport": int(
                (snapshot.get("sampling") or {}).get("historicalCount") or 0
            ),
            "summary": _sampling_summary(snapshot, inspection),
            "evidence": evidence,
            "inspection": inspection,
            "productContext": deepcopy(inspection.get("context") or {}),
            "review": deepcopy(snapshot.get("review") or {}),
            "recommendationGaps": {
                "unmappedEvidence": deepcopy(inspection.get("unmappedEvidence") or []),
                "compositionGaps": deepcopy(inspection.get("compositionGaps") or []),
                "knowledgeGaps": deepcopy(inspection.get("knowledgeGaps") or []),
            },
            "disclaimer": inspection.get("disclaimer") or SAMPLING_DISCLAIMER,
            "frozenAssets": frozen_assets,
        }

    def list_current(self) -> dict[str, Any]:
        memberships = self.sampling_store.list_current()
        items: list[dict[str, Any]] = []
        for ordinal, membership in enumerate(memberships, start=1):
            snapshot = self.data_store.get_snapshot(membership["sourceSnapshotId"])
            artifacts = build_snapshot_artifacts(
                self._run_root(snapshot), str(snapshot.get("productId") or "")
            )
            items.append(
                self._item_from_snapshot(
                    membership, snapshot, artifacts, ordinal=ordinal
                )
            )
        return {"items": items, "count": len(items)}

    def _freeze_current(
        self, list_id: str, exported_at: str
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Path]:
        snapshot_path, workbook_path = self._relative_paths(list_id)
        list_root = self._list_root(list_id)
        temporary_assets_root = list_root / f".assets-{uuid.uuid4().hex}.tmp"
        with self.data_store.transaction() as connection:
            memberships = self.sampling_store.list_current_in_transaction(connection)
            if not memberships:
                raise SamplingListEmptyError("当前抽检清单为空，无法导出")
            snapshots = [
                self.data_store.get_snapshot_in_transaction(
                    connection, membership["sourceSnapshotId"]
                )
                for membership in memberships
            ]
            if any(
                (snapshot.get("review") or {}).get("status")
                != "recommend_follow_up"
                for snapshot in snapshots
            ):
                raise SamplingExportConflictError(
                    "当前清单包含与人工复核结论不一致的商品，无法导出"
                )
            self.sampling_store.create_preparing(
                connection,
                list_id=list_id,
                snapshot_path=snapshot_path,
                workbook_path=workbook_path,
                created_at=exported_at,
            )
        created_root = False
        try:
            list_root.mkdir(parents=True, exist_ok=False)
            created_root = True
            temporary_assets_root.mkdir(parents=True)
            items: list[dict[str, Any]] = []
            for ordinal, (membership, snapshot) in enumerate(
                zip(memberships, snapshots), start=1
            ):
                artifacts = build_snapshot_artifacts(
                    self._run_root(snapshot), str(snapshot.get("productId") or "")
                )
                items.append(
                    self._item_from_snapshot(
                        membership,
                        snapshot,
                        artifacts,
                        ordinal=ordinal,
                        temporary_assets_root=temporary_assets_root,
                    )
                )
            os.replace(temporary_assets_root, list_root / "assets")
            return memberships, items, list_root
        except Exception:
            if created_root and list_root.is_dir():
                shutil.rmtree(list_root)
            raise

    def _write_frozen_files(
        self,
        list_root: Path,
        list_id: str,
        exported_at: str,
        items: list[dict[str, Any]],
    ) -> tuple[dict[str, Any], str, str]:
        workbook_path = list_root / SAMPLING_WORKBOOK_FILE
        workbook_temporary = list_root / f".{SAMPLING_WORKBOOK_FILE}.{uuid.uuid4().hex}.tmp"
        self.workbook_writer(items, list_id, exported_at, workbook_temporary)
        workbook_sha256 = file_sha256(workbook_temporary)
        os.replace(workbook_temporary, workbook_path)
        payload = {
            "schemaVersion": SAMPLING_SNAPSHOT_VERSION,
            "listId": list_id,
            "status": "exported",
            "createdAt": exported_at,
            "exportedAt": exported_at,
            "itemCount": len(items),
            "workbookFile": SAMPLING_WORKBOOK_FILE,
            "workbookSha256": workbook_sha256,
            "disclaimer": SAMPLING_DISCLAIMER,
            "items": items,
        }
        snapshot_path = list_root / SAMPLING_SNAPSHOT_FILE
        write_json(snapshot_path, payload)
        return payload, file_sha256(snapshot_path), workbook_sha256

    def _finalize_export(
        self,
        *,
        list_id: str,
        payload: dict[str, Any],
        snapshot_sha256: str,
        workbook_sha256: str,
        memberships: list[dict[str, Any]],
    ) -> dict[str, Any]:
        with self.data_store.transaction() as connection:
            metadata = self.sampling_store.finalize_export(
                connection,
                list_id=list_id,
                exported_at=str(payload["exportedAt"]),
                item_count=int(payload["itemCount"]),
                snapshot_sha256=snapshot_sha256,
                workbook_sha256=workbook_sha256,
                items=list(payload["items"]),
                frozen_memberships=memberships,
            )
        return self._metadata_urls(metadata)

    def export_current(self, *, confirmed: bool) -> dict[str, Any]:
        if confirmed is not True:
            raise SamplingExportValidationError("必须确认后才能导出当前抽检清单")
        if not self._export_lock.acquire(blocking=False):
            raise SamplingExportInProgressError("已有抽检清单正在导出")
        list_id = ""
        list_root: Path | None = None
        try:
            exported_at = iso_now()
            list_id = (
                "SL-"
                + exported_at.replace("-", "").replace(":", "").replace("+", "-")[:15]
                + "-"
                + uuid.uuid4().hex[:8]
            )
            memberships, items, list_root = self._freeze_current(
                list_id, exported_at
            )
            payload, snapshot_sha256, workbook_sha256 = self._write_frozen_files(
                list_root, list_id, exported_at, items
            )
            return self._finalize_export(
                list_id=list_id,
                payload=payload,
                snapshot_sha256=snapshot_sha256,
                workbook_sha256=workbook_sha256,
                memberships=memberships,
            )
        except Exception as exc:
            if list_id:
                self.sampling_store.delete_history(list_id)
            if list_root is not None and list_root.is_dir():
                shutil.rmtree(list_root)
            if isinstance(exc, SamplingExportError):
                raise
            raise SamplingExportError(f"抽检清单导出失败：{exc}") from exc
        finally:
            self._export_lock.release()

    def _metadata_urls(self, metadata: dict[str, Any]) -> dict[str, Any]:
        list_id = validate_list_id(metadata["listId"])
        return {
            **metadata,
            "detailUrl": f"/api/sampling-lists/{list_id}",
            "downloadUrl": f"/api/sampling-lists/{list_id}/download",
        }

    def list_history(self) -> dict[str, Any]:
        histories = [
            self._metadata_urls(item)
            for item in self.sampling_store.list_history(status="exported")
        ]
        return {"lists": histories, "count": len(histories)}

    def _load_history_payload(self, list_id: str) -> tuple[dict[str, Any], Path]:
        validated = validate_list_id(list_id)
        metadata = self.sampling_store.get_history_metadata(validated)
        if metadata is None or metadata["status"] != "exported":
            raise SamplingHistoryNotFoundError("历史抽检清单不存在")
        list_root = self._list_root(validated)
        snapshot_path = list_root / SAMPLING_SNAPSHOT_FILE
        if not snapshot_path.is_file():
            raise SamplingHistoryIntegrityError("历史清单冻结文件不存在")
        if metadata["snapshotSha256"] and file_sha256(snapshot_path) != metadata["snapshotSha256"]:
            raise SamplingHistoryIntegrityError("历史清单冻结文件校验失败")
        try:
            payload = read_json(snapshot_path)
        except (OSError, ValueError, TypeError) as exc:
            raise SamplingHistoryIntegrityError("历史清单冻结文件无法读取") from exc
        if (
            not isinstance(payload, dict)
            or payload.get("listId") != validated
            or not isinstance(payload.get("items"), list)
        ):
            raise SamplingHistoryIntegrityError("历史清单冻结文件结构不合法")
        return payload, list_root

    def get_history(self, list_id: str) -> dict[str, Any]:
        payload, _list_root = self._load_history_payload(list_id)
        result = deepcopy(payload)
        validated = validate_list_id(list_id)
        for item in result.get("items") or []:
            for asset in item.get("frozenAssets") or []:
                frozen_path = str(asset.get("frozenPath") or "")
                asset["url"] = (
                    f"/api/sampling-lists/{validated}/files/{frozen_path}"
                    if frozen_path
                    else None
                )
        result["downloadUrl"] = f"/api/sampling-lists/{validated}/download"
        return result

    def workbook_path(self, list_id: str) -> Path:
        payload, list_root = self._load_history_payload(list_id)
        workbook_path = list_root / str(
            payload.get("workbookFile") or SAMPLING_WORKBOOK_FILE
        )
        if (
            not workbook_path.resolve().is_relative_to(list_root)
            or not workbook_path.is_file()
        ):
            raise SamplingHistoryIntegrityError("历史清单工作簿不存在")
        expected = str(payload.get("workbookSha256") or "")
        if expected and file_sha256(workbook_path) != expected:
            raise SamplingHistoryIntegrityError("历史清单工作簿校验失败")
        return workbook_path

    def asset_path(self, list_id: str, relative_path: str) -> Path:
        payload, list_root = self._load_history_payload(list_id)
        normalized = str(relative_path or "").replace("\\", "/").lstrip("/")
        allowed = {
            str(asset.get("frozenPath") or "")
            for item in payload.get("items") or []
            for asset in item.get("frozenAssets") or []
        }
        if normalized not in allowed:
            raise SamplingHistoryNotFoundError("历史证据资产不存在")
        destination = (list_root / normalized).resolve()
        if not destination.is_relative_to(list_root) or not destination.is_file():
            raise SamplingHistoryNotFoundError("历史证据资产不存在")
        expected = next(
            (
                str(asset.get("sha256") or "")
                for item in payload.get("items") or []
                for asset in item.get("frozenAssets") or []
                if str(asset.get("frozenPath") or "") == normalized
            ),
            "",
        )
        if expected and file_sha256(destination) != expected:
            raise SamplingHistoryIntegrityError("历史证据资产校验失败")
        return destination

    def recover_preparing(self) -> dict[str, int]:
        recovered = 0
        cleaned = 0
        for metadata in self.sampling_store.list_history(status="preparing"):
            list_id = metadata["listId"]
            list_root = self._list_root(list_id)
            try:
                snapshot_path = list_root / SAMPLING_SNAPSHOT_FILE
                workbook_path = list_root / SAMPLING_WORKBOOK_FILE
                payload = read_json(snapshot_path)
                if (
                    not isinstance(payload, dict)
                    or payload.get("listId") != list_id
                    or payload.get("status") != "exported"
                    or not isinstance(payload.get("items"), list)
                    or not workbook_path.is_file()
                ):
                    raise SamplingHistoryIntegrityError("未完成的冻结文件不完整")
                snapshot_sha256 = file_sha256(snapshot_path)
                workbook_sha256 = file_sha256(workbook_path)
                if payload.get("workbookSha256") != workbook_sha256:
                    raise SamplingHistoryIntegrityError("未完成的工作簿校验失败")
                memberships = [
                    {
                        "productId": item["productId"],
                        "sourceSnapshotId": item["sourceSnapshotId"],
                        "sourceTaskId": item["sourceTaskId"],
                        "addedFrom": item["addedFrom"],
                        "addedAt": item["addedAt"],
                        "updatedAt": item["updatedAt"],
                    }
                    for item in payload["items"]
                ]
                self._finalize_export(
                    list_id=list_id,
                    payload=payload,
                    snapshot_sha256=snapshot_sha256,
                    workbook_sha256=workbook_sha256,
                    memberships=memberships,
                )
                recovered += 1
            except Exception:
                self.sampling_store.delete_history(list_id)
                if list_root.is_dir():
                    shutil.rmtree(list_root)
                cleaned += 1
        return {"recovered": recovered, "cleaned": cleaned}

    def rebuild_history_index(self) -> dict[str, int]:
        rebuilt = 0
        skipped = 0
        if not self.sampling_root.is_dir():
            return {"rebuilt": 0, "skipped": 0}
        for list_root in sorted(self.sampling_root.iterdir()):
            if not list_root.is_dir():
                continue
            try:
                list_id = validate_list_id(list_root.name)
                snapshot_path = list_root / SAMPLING_SNAPSHOT_FILE
                workbook_path = list_root / SAMPLING_WORKBOOK_FILE
                payload = read_json(snapshot_path)
                if (
                    not isinstance(payload, dict)
                    or payload.get("listId") != list_id
                    or payload.get("status") != "exported"
                    or not isinstance(payload.get("items"), list)
                    or not workbook_path.is_file()
                ):
                    raise SamplingHistoryIntegrityError("冻结文件不完整")
                workbook_sha256 = file_sha256(workbook_path)
                if payload.get("workbookSha256") != workbook_sha256:
                    raise SamplingHistoryIntegrityError("工作簿校验失败")
                self.sampling_store.rebuild_exported(
                    list_id=list_id,
                    exported_at=str(payload.get("exportedAt") or ""),
                    item_count=len(payload["items"]),
                    snapshot_path=f"sampling_lists/{list_id}/{SAMPLING_SNAPSHOT_FILE}",
                    workbook_path=f"sampling_lists/{list_id}/{SAMPLING_WORKBOOK_FILE}",
                    snapshot_sha256=file_sha256(snapshot_path),
                    workbook_sha256=workbook_sha256,
                    created_at=str(
                        payload.get("createdAt") or payload.get("exportedAt") or ""
                    ),
                    items=list(payload["items"]),
                )
                rebuilt += 1
            except (KeyError, OSError, TypeError, ValueError, SamplingExportError):
                skipped += 1
        return {"rebuilt": rebuilt, "skipped": skipped}
